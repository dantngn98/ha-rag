"""
Omega AI RAG - LangChain + Chroma, local bge-small-en-v1.5 embeddings, Azure AI Foundry chat model.

Builds a Chroma index over the PDF corpus, runs every evaluation question through
(1) the RAG pipeline and (2) a no-retrieval baseline using the same model, scores
retrieval automatically (precision@k, recall@k), and writes an Excel workbook with
blank columns for manual answer scoring.

Usage:
    python rag_eval.py --docs ./docs --questions omega_rag_eval_questions.xlsx
    python rag_eval.py --docs ./docs --questions omega_rag_eval_questions.xlsx --dry-run   # no API calls
"""
import argparse
import hashlib
import os
import re
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import pandas as pd
from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyMuPDFLoader
from langchain_core.documents import Document
from langchain_core.messages import HumanMessage, SystemMessage
from langchain_text_splitters import RecursiveCharacterTextSplitter
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

HANDBOOK = "Starkey_Technical_Product_Handbook.pdf"

TITLES = {
    "Starkey_Omega_AI__Wireless_Tech_Tiers_ITC_R.pdf": "Omega AI+ (Omega AI) ITC R In-The-Canal Rechargeable - technical data sheet",
    "Starkey_Omega_AI__Wireless_Tech_Tiers_RIC_312.pdf": "Omega AI+ (Omega AI) RIC 312 Receiver-In-Canal - technical data sheet",
    "Starkey_Omega_AI__Wireless_Tech_Tiers_mRIC_R_LED.pdf": "Omega AI+ (Omega AI) mRIC R micro Receiver-In-Canal Rechargeable - technical data sheet",
    "Starkey_Omega_AI__Non-Wireless_Tech_Tiers_CIC_R_NW.pdf": "Omega AI+ (Omega AI) CIC R NW Completely-In-Canal Rechargeable Non-Wireless - technical data sheet",
    "Starkey_Omega_AI__Non-Wireless_Tech_Tiers_IIC_NW.pdf": "Omega AI+ (Omega AI) IIC NW Invisible-In-Canal Non-Wireless - technical data sheet",
    HANDBOOK: "Starkey Technical Product Handbook, Fall 2026 (Omega AI+)",
    "Starkey_Accessory_Compatiblity_Guide_-_Fall_26.pdf": "Starkey Accessory Compatibility Guide, Fall 2026 (chargers, accessories, apps)",
    "DNN_360_Feature_Summary.pdf": "DNN 360 feature summary (Omega AI), Nov 2025",
    "Health_and_Wellness_Feature_Summary.pdf": "Health and Wellness feature summary (Omega AI & Edge AI), Nov 2025",
    "LE_Audio_Feature_Summary_-_O.pdf": "Bluetooth LE Audio and Auracast feature summary (Omega AI & Edge AI), Feb 2026",
}

RAG_SYSTEM = (
    "You are a technical assistant for hearing care professionals answering questions about "
    "Starkey hearing aids. Answer ONLY from the context excerpts provided. Cite each fact as "
    "[file, p. N]. If the excerpts do not contain the answer, reply exactly: "
    "'Not found in the provided documents.' Do not use outside knowledge. Be concise."
)
BASELINE_SYSTEM = (
    "You are a technical assistant for hearing care professionals answering questions about "
    "Starkey hearing aids. Answer from your own knowledge. If you do not know or are not "
    "confident, say 'I don't know.' Be concise."
)


# page handling
def printed_to_pdf_page(fname: str, printed: int) -> int:
    """The handbook PDF stores two printed pages per PDF page. Other files are 1:1."""
    return printed // 2 + 1 if fname == HANDBOOK else printed


def parse_pages(text: str) -> list[int]:
    pages = []
    for a, b, single in re.findall(r"(\d+)\s*[--]\s*(\d+)|(\d+)", text):
        pages.extend(range(int(a), int(b) + 1) if a else [int(single)])
    return pages


def relevant_set(src_cell, page_cell) -> set[tuple[str, int]]:
    """(file, pdf_page) pairs that count as a correct retrieval. Text in parentheses
    (e.g. 'trap: 21') is ignored. Empty set = unanswerable question."""
    src = "" if pd.isna(src_cell) else str(src_cell)
    if not src.strip() or src.strip().lower().startswith("none"):
        return set()
    files = [s.strip() for s in src.split(";") if s.strip().lower().endswith(".pdf")]
    page_txt = re.sub(r"\(.*?\)", "", "" if pd.isna(page_cell) else str(page_cell))
    parts = [p for p in page_txt.split(";") if p.strip()]
    rel = set()
    if len(parts) == len(files) and len(files) > 1:  # "Handbook 22; flyer 1" style
        for f, part in zip(files, parts):
            rel |= {(f, printed_to_pdf_page(f, p)) for p in parse_pages(part)}
    else:
        for f in files:
            rel |= {(f, printed_to_pdf_page(f, p)) for p in parse_pages(page_txt)}
    return rel


# models
def azure_base_url() -> str:
    ep = os.environ["AZURE_OPENAI_ENDPOINT"].rstrip("/")
    return ep + "/" if ep.endswith("/openai/v1") else ep + "/openai/v1/"


def get_embeddings(dry_run: bool):
    """Local Hugging Face embeddings."""
    if dry_run:
        from langchain_core.embeddings import DeterministicFakeEmbedding
        return DeterministicFakeEmbedding(size=256), "fake-256"
    from langchain_huggingface import HuggingFaceEmbeddings
    name = os.getenv("EMBEDDING_MODEL", "BAAI/bge-small-en-v1.5")
    query_kwargs = {"normalize_embeddings": True}
    if "bge" in name.lower() and "-en" in name.lower():
        query_kwargs["prompt"] = "Represent this sentence for searching relevant passages: "
    return HuggingFaceEmbeddings(model_name=name, encode_kwargs={"normalize_embeddings": True},
                                 query_encode_kwargs=query_kwargs), name


class LLM:
    """Wraps ChatOpenAI against the Azure v1 endpoint."""

    def __init__(self, dry_run: bool):
        self.dry_run = dry_run
        self.deployment = os.getenv("AZURE_CHAT_DEPLOYMENT", "gpt-6-luna")
        if not dry_run:
            from langchain_openai import ChatOpenAI
            self.client = ChatOpenAI(model=self.deployment, base_url=azure_base_url(),
                                     api_key=os.environ["AZURE_OPENAI_API_KEY"], max_retries=3, timeout=180)

    def ask(self, system: str, user: str) -> str:
        if self.dry_run:
            return "[dry run - no model call]"
        return self.client.invoke([SystemMessage(system), HumanMessage(user)]).content


# indexing
def load_chunks(docs_dir: Path, chunk_size: int, overlap: int, prefix: bool) -> list[Document]:
    pdfs = sorted(docs_dir.glob("*.pdf"))
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=overlap)
    chunks = []
    for pdf in pdfs:
        title = TITLES.get(pdf.name, pdf.stem.replace("_", " "))
        for page in PyMuPDFLoader(str(pdf)).load():
            text = CONTROL_CHARS.sub(" ", page.page_content).strip()  # PDFs contain stray control bytes
            if not text:
                continue
            pno = int(page.metadata.get("page", 0)) + 1  # loader is 0-indexed
            meta = {"file": pdf.name, "page": pno, "title": title}  
            for i, piece in enumerate(splitter.split_text(text)):
                content = f"[Document: {title} | page {pno}]\n{piece}" if prefix else piece
                chunks.append(Document(page_content=content, metadata={**meta, "chunk": i}))
    print(f"  {len(pdfs)} PDFs -> {len(chunks)} chunks (size={chunk_size}, overlap={overlap}, prefix={prefix})")
    return chunks


def get_vectorstore(chunks, embeddings, emb_name, docs_dir, chunk_size, overlap, prefix) -> Chroma:
    files_sig = ",".join(f"{p.name}:{p.stat().st_size}" for p in sorted(docs_dir.glob("*.pdf")))
    key = f"{emb_name}|{chunk_size}|{overlap}|{prefix}|{files_sig}"
    cfg_id = hashlib.md5(key.encode()).hexdigest()[:10]
    vs = Chroma(collection_name="omega", embedding_function=embeddings,
                persist_directory=str(Path("chroma_db") / cfg_id))
    if not vs.get(limit=1)["ids"]:
        print(f"  Embedding {len(chunks)} chunks into chroma_db/{cfg_id} …")
        ids = [f"{c.metadata['file']}|{c.metadata['page']}|{c.metadata['chunk']}" for c in chunks]
        for i in range(0, len(chunks), 64):
            vs.add_documents(chunks[i:i + 64], ids=ids[i:i + 64])
    else:
        print(f"  Reusing existing index chroma_db/{cfg_id}")
    return vs


# evaluation
def run(args):
    load_dotenv()
    docs_dir = Path(args.docs)
    overlap = args.overlap if args.overlap is not None else int(args.chunk_size * 0.15)

    q = pd.read_excel(args.questions, sheet_name="Questions").dropna(subset=["ID", "Question"])
    print(f"Loaded {len(q)} questions")

    embeddings, emb_name = get_embeddings(args.dry_run)
    chunks = load_chunks(docs_dir, args.chunk_size, overlap, args.prefix)
    vs = get_vectorstore(chunks, embeddings, emb_name, docs_dir, args.chunk_size, overlap, args.prefix)
    llm = LLM(args.dry_run)
    # How many chunks each (file, page) produced - needed for the precision ceiling.
    page_chunks = Counter((c.metadata["file"], int(c.metadata["page"])) for c in chunks)

    results, retrieved_rows = [], []
    for _, row in q.iterrows():
        qid, question = str(row["ID"]), str(row["Question"])
        print(f"[{qid}] {question}")
        rel = relevant_set(row.get("Source document"), row.get("Page"))

        hits = vs.similarity_search_with_score(question, k=args.k)
        first_rank, file_hit = None, False
        rel_chunks_hit, rel_pages_hit = 0, set()
        context_blocks = []
        for rank, (doc, dist) in enumerate(hits, start=1):
            f, p = doc.metadata["file"], int(doc.metadata["page"])
            is_rel = (f, p) in rel
            file_hit |= any(f == rf for rf, _ in rel)
            if is_rel:
                rel_chunks_hit += 1
                rel_pages_hit.add((f, p))
                if first_rank is None:
                    first_rank = rank
            context_blocks.append(f"[{rank}] ({f}, p. {p})\n{doc.page_content}")
            retrieved_rows.append({"ID": qid, "Rank": rank, "File": f, "PDF page": p,
                                   "Distance": round(float(dist), 4), "Relevant": "Y" if is_rel else "",
                                   "Chunk text": doc.page_content})

        context = "\n\n".join(context_blocks)
        t0 = time.time()
        rag_answer = llm.ask(RAG_SYSTEM, f"Context excerpts:\n\n{context}\n\nQuestion: {question}")
        rag_s = time.time() - t0
        base_answer = "" if args.skip_baseline else llm.ask(BASELINE_SYSTEM, question)

        answerable = bool(rel)
        results.append({
            "ID": qid, "Question": question, "Expected answer": row.get("Expected answer", ""),
            "RAG answer": rag_answer, "Baseline answer": base_answer,
            "Retrieved (rank: file p.N)": "\n".join(f"{r}: {d.metadata['file']} p.{d.metadata['page']}"
                                                   for r, (d, _) in enumerate(hits, 1)),
            # Precision@k: share of the k retrieved chunks that come from a relevant page.
            f"Precision@{args.k}": round(rel_chunks_hit / len(hits), 3) if answerable else None,
            # Recall@k: share of the relevant pages that appear anywhere in the top k.
            f"Recall@{args.k}": round(len(rel_pages_hit) / len(rel), 3) if answerable else None,
            # Best precision possible: a relevant page that produced only 1 chunk caps precision@4 at 0.25.
            "Precision ceiling": round(min(len(hits), sum(page_chunks[x] for x in rel)) / len(hits), 3) if answerable else None,
            "Right file retrieved": (1 if file_hit else 0) if answerable else None,
            "RAG latency (s)": round(rag_s, 1),
            "RAG correct (0/1/2)": None, "RAG faithful (Y/N)": None,
            "Baseline correct (0/1/2)": None, "Baseline hallucinated (Y/N)": None, "Notes": None,
        })

    config = {
        "Run time": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Chat deployment": llm.deployment if not args.dry_run else "dry run",
        "Embeddings": emb_name, "Chunk size": args.chunk_size, "Chunk overlap": overlap,
        "Title prefix on chunks": args.prefix, "Top k": args.k, "Chunks indexed": len(chunks),
        "PDFs": ", ".join(p.name for p in sorted(docs_dir.glob("*.pdf"))),
        "Baseline run": not args.skip_baseline,
    }
    out = Path(args.out or f"results_cs{args.chunk_size}_{'prefix' if args.prefix else 'noprefix'}_k{args.k}"
               f"{'_DRYRUN' if args.dry_run else ''}.xlsx")
    write_workbook(out, results, retrieved_rows, config, args.k)
    print(f"\nWrote {out}")


def write_workbook(out, results, retrieved_rows, config, k):
    clean = lambda df: df.apply(lambda c: c.map(lambda v: ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v))
    res = clean(pd.DataFrame(results))
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        res.to_excel(xw, sheet_name="Results", index=False)
        clean(pd.DataFrame(retrieved_rows)).to_excel(xw, sheet_name="Retrieved chunks", index=False)
        pd.DataFrame(list(config.items()), columns=["Setting", "Value"]).to_excel(xw, sheet_name="Run config", index=False)

    wb = load_workbook(out)
    ws = wb["Results"]
    n = ws.max_row
    col = {c.value: c.column_letter for c in ws[1]}
    manual = ["RAG correct (0/1/2)", "RAG faithful (Y/N)", "Baseline correct (0/1/2)",
              "Baseline hallucinated (Y/N)", "Notes"]
    widths = {"ID": 6, "Question": 40, "Expected answer": 40, "RAG answer": 55, "Baseline answer": 55,
              "Retrieved (rank: file p.N)": 50}
    yellow = PatternFill("solid", start_color="FFF2CC")
    for name, letter in col.items():
        ws.column_dimensions[letter].width = widths.get(name, 14)
    for r in ws.iter_rows(min_row=1, max_row=n):
        for c in r:
            c.font = Font(name="Arial", size=10, bold=(c.row == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.row > 1 and ws.cell(1, c.column).value in manual:
                c.fill = yellow
    ws.freeze_panes = "C2"

    rc = wb["Retrieved chunks"]
    for letter, w in zip("ABCDEFG", [6, 6, 45, 9, 10, 9, 90]):
        rc.column_dimensions[letter].width = w
    for r in rc.iter_rows():
        for c in r:
            c.font = Font(name="Arial", size=9, bold=(c.row == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb["Run config"].column_dimensions["A"].width = 24
    wb["Run config"].column_dimensions["B"].width = 100

    s = wb.create_sheet("Summary", 0)
    rng = lambda name: f"Results!${col[name]}$2:${col[name]}${n}"
    rows = [
        ("Retrieval (answerable questions only)", None),
        (f"Mean precision@{k}", f"=IFERROR(AVERAGE({rng(f'Precision@{k}')}),\"\")"),
        (f"Mean precision ceiling (best possible)", f"=IFERROR(AVERAGE({rng('Precision ceiling')}),\"\")"),
        (f"Mean recall@{k}", f"=IFERROR(AVERAGE({rng(f'Recall@{k}')}),\"\")"),
        ("Right file retrieved", f"=IFERROR(AVERAGE({rng('Right file retrieved')}),\"\")"),
        ("Answers (fill yellow columns on Results)", None),
        ("RAG mean correctness (0-2)", f"=IFERROR(AVERAGE({rng('RAG correct (0/1/2)')}),\"\")"),
        ("RAG fully correct (count of 2s)", f"=COUNTIF({rng('RAG correct (0/1/2)')},2)"),
        ("Baseline mean correctness (0-2)", f"=IFERROR(AVERAGE({rng('Baseline correct (0/1/2)')}),\"\")"),
        ("Baseline fully correct (count of 2s)", f"=COUNTIF({rng('Baseline correct (0/1/2)')},2)"),
        ("RAG unfaithful answers", f"=COUNTIF({rng('RAG faithful (Y/N)')},\"N\")"),
        ("Baseline hallucinations", f"=COUNTIF({rng('Baseline hallucinated (Y/N)')},\"Y\")"),
    ]
    for i, (label, formula) in enumerate(rows, start=1):
        s.cell(i, 1, label).font = Font(name="Arial", bold=formula is None)
        if formula:
            c = s.cell(i, 2, formula)
            c.font = Font(name="Arial")
            c.number_format = "0" if "COUNTIF" in formula else "0.00"
    s.column_dimensions["A"].width = 42
    s.column_dimensions["B"].width = 12
    s.cell(len(rows) + 2, 1, "Scoring: correct 2 = fully right, 1 = partially right, 0 = wrong. "
                  "Faithful = every claim supported by the retrieved chunks. "
                  "Hallucinated = a confident specific claim that is wrong.").font = Font(name="Arial", italic=True, size=9)
    wb.save(out)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--docs", required=True, help="Folder with the corpus PDFs")
    ap.add_argument("--questions", required=True, help="Eval questions .xlsx (sheet 'Questions')")
    ap.add_argument("--chunk-size", type=int, default=800)
    ap.add_argument("--overlap", type=int, default=None, help="Default: 15%% of chunk size")
    ap.add_argument("--k", type=int, default=4)
    ap.add_argument("--no-prefix", dest="prefix", action="store_false", help="Ablation: no title header on chunks")
    ap.add_argument("--skip-baseline", action="store_true", help="Skip no-retrieval baseline (saves calls on ablation runs)")
    ap.add_argument("--dry-run", action="store_true", help="Fake embeddings and no LLM calls - tests plumbing only")
    ap.add_argument("--out", default=None)
    run(ap.parse_args())
